import json
from docx import Document
from pptx import Presentation
from app.document_processing import clean_text, chunk_sections, extract_sections
from app.ingest import load_manifest, parse_actions
import pytest


def test_word_text_and_tables_have_locations(tmp_path):
    path = tmp_path / "notes.docx"
    doc = Document()
    doc.add_heading("LangChain", level=1)
    doc.add_paragraph("LangChain connects models and tools.")
    doc.add_table(rows=1, cols=2).cell(0, 0).text = "Alembic"
    doc.save(path)
    result = extract_sections(path)
    assert any("LangChain connects models" in section.text for section in result)
    assert any(section.location == "table 1" and "Alembic" in section.text for section in result)


def test_powerpoint_retains_slide_location(tmp_path):
    path = tmp_path / "slides.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "RAG"
    slide.placeholders[1].text = "Retrieve evidence before answering."
    deck.save(path)
    result = extract_sections(path)
    assert result[0].location == "slide 1"
    assert "Retrieve evidence" in result[0].text


def test_markdown_splits_at_headings(tmp_path):
    path = tmp_path / "meeting.md"
    path.write_text("# Decisions\nUse pgvector.\n## Actions\nDocument ingestion.", encoding="utf-8")

    sections = extract_sections(path)

    assert [section.strategy for section in sections] == ["markdown-heading", "markdown-heading"]
    assert "Decisions" in sections[0].location
    assert "Actions" in sections[1].location


def test_python_keeps_functions_as_sections(tmp_path):
    path = tmp_path / "example.py"
    path.write_text(
        "import os\n\n@staticmethod\ndef connect():\n    return os.getenv('URL')\n\nDEFAULT_PORT = 8000\n",
        encoding="utf-8",
    )

    sections = extract_sections(path)

    assert any("function connect" in section.location for section in sections)
    assert any("@staticmethod" in section.text for section in sections)
    assert any("DEFAULT_PORT" in section.text for section in sections)


def test_large_section_is_split_but_keeps_location():
    from app.document_processing import ExtractedSection

    chunks = chunk_sections([ExtractedSection("page 3", "word " * 1000, "pdf-page")])

    assert len(chunks) > 1
    assert all(chunk.metadata["location"].startswith("page 3") for chunk in chunks)
    assert all(chunk.metadata["strategy"] == "pdf-page" for chunk in chunks)


def test_sql_splits_at_statement_boundaries(tmp_path):
    path = tmp_path / "migration.sql"
    path.write_text("CREATE TABLE notes (id INT);\nDROP TABLE notes;", encoding="utf-8")

    sections = extract_sections(path)

    assert len(sections) == 2
    assert all(section.strategy == "sql-statement" for section in sections)


def test_csv_header_only_is_still_searchable(tmp_path):
    path = tmp_path / "index.csv"
    path.write_text("title,topic\n", encoding="utf-8")

    sections = extract_sections(path)

    assert len(sections) == 1
    assert sections[0].text == "title | topic"


def test_manifest_fails_on_missing_file(tmp_path):
    manifest = tmp_path / "sources.json"
    manifest.write_text(json.dumps([dict(id="a", path="missing.md", title="Missing", source_type="google_drive", source_url="https://example.com/missing")]))
    with pytest.raises(ValueError, match="missing"):
        load_manifest(manifest)


def test_action_table_tracks_supported_statuses():
    meeting_note = """\
| Action | Owner | Status | Due |
| --- | --- | --- | --- |
| Review migration | Tri | open | 2026-10-09 |
| Validate citations | Hieu | in_progress | 2026-10-10 |
| Publish notes | Team | done | |
"""
    actions = parse_actions(meeting_note)
    assert len(actions) == 3
    assert {a["status"] for a in actions} == {"open", "in_progress", "done"}


def test_clean_text_normalizes_malformed_pdf_text():
    assert clean_text("model\x00 size \ud835") == "model size ?"
