"""Turn downloaded files into useful, citation-friendly sections and chunks.

Different files communicate structure in different ways. A meeting note uses
headings, a presentation uses slides, and Python uses functions and classes.
Splitting every file after the same number of characters can separate a heading
from its explanation or a function signature from its body.

This module therefore has two small jobs:

1. Extract natural sections for each supported file type.
2. Split only sections that are too large for a useful embedding.

No chat model is called here. The rules are local, repeatable, and inexpensive.
"""

import ast
import csv
import re
from dataclasses import dataclass
from pathlib import Path

from langchain_core.documents import Document as LangChainDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter


# These are character limits rather than model-token limits. This keeps the
# implementation independent of a hosted tokenizer and works conservatively
# with the local EmbeddingGemma model.
MAX_CHUNK_CHARACTERS = 1_800
CHUNK_OVERLAP_CHARACTERS = 150


@dataclass(frozen=True)
class ExtractedSection:
    """One meaningful section with the location shown in a citation."""

    location: str
    text: str
    strategy: str


def clean_text(value: str) -> str:
    """Remove malformed characters that Ollama or PostgreSQL cannot accept."""

    without_nul = value.replace("\x00", "")
    return without_nul.encode("utf-8", errors="replace").decode("utf-8")


def _section(location: str, text: str, strategy: str) -> ExtractedSection:
    """Create a cleaned section so every extractor follows the same rule."""

    return ExtractedSection(location, clean_text(text).strip(), strategy)


def _extract_pdf(path: Path) -> list[ExtractedSection]:
    """Extract one section per PDF page using Trang's PyMuPDF approach."""

    import pymupdf

    sections = []
    with pymupdf.open(path) as pdf:
        for page_number, page in enumerate(pdf, start=1):
            # sort=True produces a more natural reading order for many PDFs.
            sections.append(
                _section(f"page {page_number}", page.get_text("text", sort=True), "pdf-page")
            )
    return sections


def _extract_word(path: Path) -> list[ExtractedSection]:
    """Group Word paragraphs under their heading and keep tables intact."""

    from docx import Document

    document = Document(path)
    sections = []
    heading = "Document start"
    paragraph_start = 1
    paragraph_text = []

    def save_paragraph_section(paragraph_end: int) -> None:
        if paragraph_text:
            location = f"paragraphs {paragraph_start}-{paragraph_end}: {heading}"
            sections.append(_section(location, "\n".join(paragraph_text), "word-heading"))

    for number, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if not text:
            continue
        is_heading = paragraph.style and paragraph.style.name.lower().startswith("heading")
        if is_heading:
            save_paragraph_section(number - 1)
            heading = text
            paragraph_start = number
            paragraph_text = [text]
        else:
            paragraph_text.append(text)

    save_paragraph_section(len(document.paragraphs))

    for table_number, table in enumerate(document.tables, start=1):
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        sections.append(_section(f"table {table_number}", "\n".join(rows), "word-table"))
    return sections


def _extract_powerpoint(path: Path) -> list[ExtractedSection]:
    """Keep each slide's title, body, tables, and speaker notes together."""

    from pptx import Presentation

    sections = []
    for slide_number, slide in enumerate(Presentation(path).slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.strip():
                parts.append(shape.text.strip())
            if shape.has_table:
                rows = [
                    " | ".join(cell.text.strip() for cell in row.cells)
                    for row in shape.table.rows
                ]
                parts.append("\n".join(rows))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Speaker notes:\n{notes}")
        sections.append(_section(f"slide {slide_number}", "\n\n".join(parts), "slide"))
    return sections


def _extract_markdown(path: Path) -> list[ExtractedSection]:
    """Split Markdown at headings while retaining exact source line ranges."""

    lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
    heading_pattern = re.compile(r"^#{1,6}\s+(.+)$")
    starts = [index for index, line in enumerate(lines) if heading_pattern.match(line)]
    if not starts or starts[0] != 0:
        starts.insert(0, 0)
    starts.append(len(lines))

    sections = []
    for start, end in zip(starts, starts[1:]):
        text = "\n".join(lines[start:end]).strip()
        if not text:
            continue
        match = heading_pattern.match(lines[start]) if start < len(lines) else None
        label = match.group(1) if match else "Document start"
        sections.append(_section(f"lines {start + 1}-{end}: {label}", text, "markdown-heading"))
    return sections


def _extract_python(path: Path) -> list[ExtractedSection]:
    """Keep top-level Python functions and classes as meaningful sections."""

    source = path.read_text(encoding="utf-8", errors="replace")
    lines = source.splitlines()
    try:
        tree = ast.parse(source)
    except SyntaxError:
        return [_section(f"lines 1-{len(lines)}", source, "python-fallback")]

    nodes = [
        node
        for node in tree.body
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef))
    ]
    if not nodes:
        return [_section(f"lines 1-{len(lines)}", source, "python-module")]

    sections = []
    next_unclaimed_line = 1
    for node in nodes:
        decorator_lines = [decorator.lineno for decorator in node.decorator_list]
        start_line = min([node.lineno, *decorator_lines])
        end_line = getattr(node, "end_lineno", node.lineno)
        if next_unclaimed_line < start_line:
            module_text = "\n".join(lines[next_unclaimed_line - 1 : start_line - 1])
            sections.append(
                _section(
                    f"lines {next_unclaimed_line}-{start_line - 1}: module code",
                    module_text,
                    "python-module",
                )
            )
        label = f"{type(node).__name__.removesuffix('Def').lower()} {node.name}"
        text = "\n".join(lines[start_line - 1 : end_line])
        sections.append(_section(f"lines {start_line}-{end_line}: {label}", text, "python-symbol"))
        next_unclaimed_line = end_line + 1

    if next_unclaimed_line <= len(lines):
        sections.append(
            _section(
                f"lines {next_unclaimed_line}-{len(lines)}: module code",
                "\n".join(lines[next_unclaimed_line - 1 :]),
                "python-module",
            )
        )
    return sections


def _extract_sql(path: Path) -> list[ExtractedSection]:
    """Group a simple SQL script by semicolon-terminated statements."""

    lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
    sections = []
    start = 1
    current = []
    for number, line in enumerate(lines, start=1):
        current.append(line)
        if ";" in line:
            sections.append(_section(f"lines {start}-{number}", "\n".join(current), "sql-statement"))
            start = number + 1
            current = []
    if current:
        sections.append(_section(f"lines {start}-{len(lines)}", "\n".join(current), "sql-statement"))
    return sections


def _extract_csv(path: Path) -> list[ExtractedSection]:
    """Create readable row batches and repeat the header in every batch."""

    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        rows = list(csv.reader(handle))
    if not rows:
        return []

    header, data_rows = rows[0], rows[1:]
    if not data_rows:
        return [_section("header row", " | ".join(header), "csv-rows")]

    sections = []
    batch_size = 25
    for offset in range(0, len(data_rows), batch_size):
        batch = data_rows[offset : offset + batch_size]
        text_rows = [header, *batch]
        text = "\n".join(" | ".join(row) for row in text_rows)
        first_row = offset + 2
        last_row = offset + len(batch) + 1
        sections.append(_section(f"rows {first_row}-{last_row}", text, "csv-rows"))
    return sections


def _extract_plain_text(path: Path) -> list[ExtractedSection]:
    """Return plain text as one section; the size limiter handles long files."""

    text = path.read_text(encoding="utf-8", errors="replace")
    line_count = len(text.splitlines())
    return [_section(f"lines 1-{line_count}", text, "plain-text")]


def extract_sections(path: Path) -> list[ExtractedSection]:
    """Choose the extractor that understands this file's natural structure."""

    extractors = {
        ".pdf": _extract_pdf,
        ".docx": _extract_word,
        ".pptx": _extract_powerpoint,
        ".md": _extract_markdown,
        ".py": _extract_python,
        ".sql": _extract_sql,
        ".csv": _extract_csv,
        ".txt": _extract_plain_text,
    }
    extractor = extractors.get(path.suffix.lower())
    if extractor is None:
        raise ValueError(f"Unsupported file extension: {path.suffix.lower()}")
    return extractor(path)


def chunk_sections(sections: list[ExtractedSection]) -> list[LangChainDocument]:
    """Keep small sections whole and split only sections that exceed the limit."""

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=MAX_CHUNK_CHARACTERS,
        chunk_overlap=CHUNK_OVERLAP_CHARACTERS,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    chunks = []
    for section in sections:
        if not section.text:
            continue
        pieces = splitter.split_text(section.text)
        for part_number, piece in enumerate(pieces, start=1):
            location = section.location
            if len(pieces) > 1:
                location = f"{location}, part {part_number}"
            chunks.append(
                LangChainDocument(
                    page_content=piece,
                    metadata={"location": location, "strategy": section.strategy},
                )
            )
    return chunks
