import io
import json
import logging
from pathlib import Path

from langchain_core.documents import Document

MAX_SNIFF_BYTES = 5_000_000

logging.getLogger("pypdf").setLevel(logging.ERROR)


def _src(path: Path, data: bytes | None):
    """Where to read from: the file on disk, or bytes already in memory (used by the Drive API mode)."""
    return io.BytesIO(data) if data is not None else str(path)


def _clean(text: str) -> str:
    """PDFs sometimes give half of a UTF-16 surrogate pair (math letters like the italic x).
    Re-join the valid pairs and replace broken halves, so the text can be saved as UTF-8."""
    return text.encode("utf-16", "surrogatepass").decode("utf-16", "replace")


def _doc(text: str, path: Path, page) -> Document:
    return Document(
        page_content=_clean(text),
        metadata={"source": path.name, "file_type": path.suffix.lower().lstrip(".") or "text", "page": page},
    )


def load_pdf(path: Path, data: bytes | None = None) -> list[Document]:
    from pypdf import PdfReader

    docs = []
    for i, page in enumerate(PdfReader(_src(path, data)).pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:  
            docs.append(_doc(text, path, i))
    return docs


def load_docx(path: Path, data: bytes | None = None) -> list[Document]:
    import docx

    d = docx.Document(_src(path, data))
    lines = []
    for p in d.paragraphs:
        if not p.text.strip():
            continue
        style = (p.style.name if p.style is not None else "") or ""  
        if style.startswith("Heading"):
            level = "".join(c for c in style if c.isdigit()) or "1"
            lines.append("#" * int(level) + " " + p.text.strip())
        else:
            lines.append(p.text.strip())
    for t in d.tables:
        for row in t.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return [_doc("\n".join(lines), path, 1)] if lines else []


def _table_to_text(df) -> str:
    """Each row becomes one line:  column: value | column: value"""
    df = df.dropna(how="all").fillna("")
    rows = []
    for _, row in df.iterrows():
        rows.append(" | ".join(f"{col}: {val}" for col, val in row.items() if str(val).strip()))
    return "\n".join(rows)


def load_excel(path: Path, data: bytes | None = None) -> list[Document]:
    import pandas as pd

    docs = []
    for sheet, df in pd.read_excel(_src(path, data), sheet_name=None).items():
        text = _table_to_text(df)
        if text:
            docs.append(_doc(text, path, sheet))
    return docs


def load_csv(path: Path, data: bytes | None = None) -> list[Document]:
    import pandas as pd

    text = _table_to_text(pd.read_csv(_src(path, data)))
    return [_doc(text, path, 1)] if text else []


def load_text(path: Path, data: bytes | None = None) -> list[Document]:
    if data is not None:
        text = data.decode("utf-8", errors="ignore").strip()
    else:
        text = path.read_text(encoding="utf-8", errors="ignore").strip()
    return [_doc(text, path, 1)] if text else []


def load_pptx(path: Path, data: bytes | None = None) -> list[Document]:
    from pptx import Presentation

    docs = []
    for i, slide in enumerate(Presentation(_src(path, data)).slides, start=1):
        parts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        if parts:
            docs.append(_doc("\n".join(parts), path, i))
    return docs


def load_ipynb(path: Path, data: bytes | None = None) -> list[Document]:
    """Jupyter notebook: keep the text of markdown and code cells, drop the outputs (often huge images)."""
    raw = data if data is not None else path.read_bytes()
    parts = []
    for cell in json.loads(raw.decode("utf-8", errors="ignore")).get("cells", []):
        src = "".join(cell.get("source", [])).strip()
        if src:
            parts.append(src if cell.get("cell_type") == "markdown" else "```python\n" + src + "\n```")
    return [_doc("\n\n".join(parts), path, 1)] if parts else []


TEXT_EXTS = [".md", ".txt", ".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".c", ".cpp", ".h", ".go", ".rs",
             ".sh", ".sql", ".html", ".htm", ".css", ".xml", ".yaml", ".yml", ".toml", ".ini", ".cfg",
             ".rst", ".log", ".tex", ".json", ".jsonl"]

LOADERS = {
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".xlsx": load_excel,
    ".xls": load_excel,
    ".csv": load_csv,
    ".pptx": load_pptx,
    ".ipynb": load_ipynb,
    **{ext: load_text for ext in TEXT_EXTS},
}
SUPPORTED = set(LOADERS)

BINARY_EXTS = {".db", ".sqlite", ".sqlite3", ".pyc", ".pkl", ".pickle", ".parquet", ".npy", ".npz",
               ".exe", ".dll", ".so", ".bin", ".zip", ".gz", ".tar", ".7z", ".rar", ".jar",
               ".jpg", ".jpeg", ".png", ".gif", ".bmp", ".ico", ".webp", ".svgz",
               ".mp3", ".mp4", ".wav", ".mov", ".avi", ".woff", ".woff2", ".ttf", ".otf"}


class SkipFile(Exception):
    """The file is not readable text (binary, or too big to guess): report it as SKIPPED."""


def load_file(path: Path, data: bytes | None = None) -> list[Document]:
    """data=None: read `path` from disk. data=bytes: read that content (path is only used for the name)."""
    return LOADERS[path.suffix.lower()](path, data)


def load_any(path: Path, data: bytes | None = None) -> list[Document]:
    """Known type -> its loader. Unknown type (no extension, odd extension) -> read as text
    if it looks like text; raise SkipFile for binaries."""
    if path.suffix.lower() in LOADERS:
        return load_file(path, data)
    if path.suffix.lower() in BINARY_EXTS:
        raise SkipFile("binary file")
    if data is None and path.stat().st_size > MAX_SNIFF_BYTES:
        raise SkipFile("too large to treat as text")
    raw = data if data is not None else path.read_bytes()
    if len(raw) > MAX_SNIFF_BYTES:
        raise SkipFile("too large to treat as text")
    if b"\x00" in raw[:8192]:
        raise SkipFile("binary file")
    return load_text(path, raw)


def load_folder(folder: Path):
    """Load every readable file. Returns (documents, report) where report has one row per file."""
    docs, report = [], []
    for path in sorted(p for p in folder.rglob("*") if p.is_file()):
        if any(part.startswith(".") for part in path.relative_to(folder).parts):
            continue  
        try:
            file_docs = load_any(path)
            status = "ok" if file_docs else "EMPTY (scanned PDF? needs OCR)"
        except SkipFile as e:
            file_docs, status = [], f"SKIPPED ({e})"
        except Exception as e:  
            file_docs, status = [], f"ERROR: {type(e).__name__}: {e}"
        docs.extend(file_docs)
        report.append({
            "file": path.name,
            "status": status,
            "parts": len(file_docs),
            "chars": sum(len(d.page_content) for d in file_docs),
        })
    return docs, report


def save_docs(docs: list[Document], out_file: Path) -> None:
    out_file.parent.mkdir(parents=True, exist_ok=True)
    with out_file.open("w", encoding="utf-8") as f:
        for d in docs:
            f.write(json.dumps({"text": d.page_content, "metadata": d.metadata}, ensure_ascii=False) + "\n")


def read_docs(in_file: Path) -> list[Document]:
    if not in_file.exists():
        raise SystemExit(f"{in_file} not found. Run phase1_ingest.py first.")
    with in_file.open(encoding="utf-8") as f:
        return [Document(page_content=r["text"], metadata=r["metadata"]) for r in map(json.loads, f)]
